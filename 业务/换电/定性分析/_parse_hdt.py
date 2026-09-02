import openpyxl, sys
path = r'd:\AI\证券投资\投研\能源\宁德时代\业务\换电\定性分析\重卡HDT换电规模测算.xlsx'
wb = openpyxl.load_workbook(path, data_only=False)
print('SHEETS:', wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print('==='*3, s, ws.max_row, 'x', ws.max_column, '==='*3)
    for r in range(1, ws.max_row+1):
        row_cells = []
        for c in range(1, ws.max_column+1):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v is None:
                continue
            if isinstance(v, str) and v.startswith('='):
                row_cells.append(f"{cell.coordinate}={v}")
            else:
                row_cells.append(f"{cell.coordinate}:{v}")
        if row_cells:
            print(' | '.join(row_cells))
