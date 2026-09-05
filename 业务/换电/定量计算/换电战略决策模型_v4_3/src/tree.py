"""「以终为始」决策树：根＝要做的判断，叶＝base.toml 里的一个参数。

它取代了什么
------------
原 `chain_table.py` 是一张**扁平行表**：层级靠 `↳ / ↳↳` 字符串前缀表示，按模型的
计算顺序（车辆→运营→站数→电池→收入→估值）自底向上排列，每行可选挂一个独立的
check，没挂的静默错下去。它的职责已全部并入本模块（存档见 `_archive/chain_table.py`）。

本模块反过来：

* **根是决策，不是车辆。** 自顶向下分解，任何一个结论都能一路点到参数。
* **树是真的递归结构。** 可折叠、可从任意节点下钻、可导出任意格式。
* **每个非叶节点都必须自洽。** 不变量：

      combine(子节点的模型值) == 本节点的模型值

  对不上就在**那个节点**当场报出来，并指出差多少。模型改了公式而树没跟上，
  重跑立刻暴露，不会悄悄给出一张错表。
* **每个叶节点必须有信源。** 信源直接取 base.toml 里该参数上方的注释
  （单一事实源，不另立台账）；取不到的在审计里单独列为"信源缺失"。

三情景
------
中性档严格等于模型基线（不改任何参数即基线）；悲观／乐观档同时拨动
服务费、充电段份额、私家车渗透、制造净利率四个本来就有分档的参数。
每个节点都带三档值，所以「同一指标在不同情景下差多少」不需要另一个程序。

    （本模块合并了原 chain_table.py 的职责：情景对比归入节点的三档值，
      车型分组归入节点的子节点——分组本来就是树的自然结构。）

运行
----
    python src/tree.py            # 跑全量校验，打印审计报告
    python src/tree.py --json     # 额外写出 outputs/tree.json
    python src/tree.py --xlsx     # 额外写出 outputs/换电决策树_v4.3.xlsx（可折叠层级）
"""
from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SRC = Path(__file__).resolve().parent
ROOT = SRC.parent
sys.path.insert(0, str(SRC))

from config_loader import DEFAULT_CONFIG, load_config  # noqa: E402
from lab import load_param_docs  # noqa: E402
from model import build_model  # noqa: E402
from schemas import ModelSnapshot  # noqa: E402

TREE_PATH = ROOT / "outputs" / "tree.json"
XLSX_PATH = ROOT / "outputs" / "换电决策树_v4.3.xlsx"

# 三情景 = 把 config 里本来就有三档的参数同时拨档；中性档严格等于模型基线。
SCENARIOS: list[tuple[str, dict]] = [
    ("悲观", dict(service_fee=0.20, charge_share="悲观", private="保守", margin=0.125)),
    ("中性", dict(service_fee=0.30, charge_share="中性", private="中枢", margin=0.135)),
    ("乐观", dict(service_fee=0.40, charge_share="乐观", private="激进", margin=0.150)),
]


def load_trailing_comments(path: Path = DEFAULT_CONFIG) -> dict[str, str]:
    """行尾注释也算信源：`rte = 0.92   # 站效率 ...`。

    lab.load_param_docs 只认键上方的注释块，行尾注释会被漏掉，
    据此判"缺信源"会误伤。两者并用才是完整的信源视图。
    """
    import re as _re
    out: dict[str, str] = {}
    stack: list[str] = []
    key_re = _re.compile(r"^\s*(?:\[+([^\]]+)\]+|([A-Za-z_][\w.]*)\s*=)")
    for raw in path.read_text("utf-8").splitlines():
        m = key_re.match(raw)
        if not m:
            continue
        if m.group(1):
            stack = [m.group(1).strip().strip("[]").strip()]
            continue
        if "#" not in raw:
            continue
        comment = raw.split("#", 1)[1].strip()
        if comment:
            out[".".join(stack + [m.group(2)])] = comment
    return out


# ─────────────────────────────────────────── 上下文
@dataclass
class Ctx:
    cfg: dict
    snap: dict          # asdict(ModelSnapshot)
    docs: dict          # base.toml 键上方注释块 → {"reason": ..., "section": ...}
    tail: dict          # base.toml 行尾注释 → 说明文字

    def src_of(self, param: str) -> str:
        """参数的信源说明：上方注释块优先，其次行尾注释。"""
        reason = (self.docs.get(param) or {}).get("reason", "").strip()
        return reason or self.tail.get(param, "").strip()

    def p(self, path: str) -> Any:
        """读 base.toml 参数。"""
        node: Any = self.cfg
        for key in path.split("."):
            node = node[key]
        return node

    def m(self, path: str) -> Any:
        """读模型实跑值（快照路径）。"""
        node: Any = self.snap
        for key in path.split("."):
            node = node[key]
        return node


def build_ctx(scen: dict | None = None) -> Ctx:
    """构造一个情景的上下文；scen=None 即基线（＝中性档）。"""
    from dataclasses import asdict
    cfg = load_config()
    if scen:
        cfg["swap_business"]["service_fee_rmb_kwh"] = scen["service_fee"]
        cfg["charge_share"]["scenario"] = scen["charge_share"]
        cfg["swap_business"]["with_swap_manufacturing_net_margin"] = scen["margin"]
    snap = build_model(cfg, private_scenario=(scen or {}).get("private", "中枢"))
    return Ctx(cfg=cfg, snap=asdict(snap), docs=load_param_docs(),
               tail=load_trailing_comments())


# ─────────────────────────────────────────── 节点
@dataclass
class Node:
    key: str
    label: str
    unit: str
    value: Callable[[Ctx], float]
    formula: str = ""                                  # 中文名运算，绝不写数字
    combine: Callable[..., float] | None = None        # 用子节点值重算本节点
    children: list["Node"] = field(default_factory=list)
    param: str | None = None                           # 叶节点：base.toml 路径
    source: str | None = None                          # 叶节点：信源覆盖（默认取 toml 注释）
    tol: float = 1e-9                                  # 相对容差
    tol_note: str = ""                                 # 放宽容差的理由（必须写）

    @property
    def is_leaf(self) -> bool:
        return not self.children


def N(key, label, unit, value, formula="", combine=None, children=None,
      param=None, source=None, tol=1e-9, tol_note="") -> Node:
    return Node(key, label, unit, value, formula, combine, children or [],
                param, source, tol, tol_note)


def P(key, label, unit, path, tol_note="") -> Node:
    """叶节点：直接是 base.toml 里的一个参数。"""
    return N(key, label, unit, lambda c, _p=path: float(c.p(_p)), param=path)


def _body_total(c: "Ctx") -> float:
    """站体总初装（建设期排期 ＋ 期初存量站）。"""
    return (
        sum(r["station_body_capex_yi"] for r in c.m("capex.annual"))
        + sum(
            c.m("capex.opening_station_stock_by_pool")[pk]
            * c.p(f"stations.{pk}.station_body_capex_wan") / 1e4
            for pk in c.m("capex.opening_station_stock_by_pool")
        )
    )


# ─────────────────────────────────────────── 树
def build_tree(c: Ctx) -> Node:
    cfg, m = c.p, c.m
    POOLS = list(m("scale.target_station_demand").keys())

    # ── 运营层：收入 ───────────────────────────
    energy = N(
        "ops.energy", "年换电交易电量（成熟期）", "亿kWh",
        lambda c: c.m("swap_business.annual_energy_yi_kwh"),
        "Σ_池( 日换电次数 × 池内加权装车电量 × 可用电量系数 × 运营天数 )",
        combine=lambda daily, onboard_avg, usable, days: (
            daily * onboard_avg * usable * days / 1e8
        ),
        children=[
            N("ops.daily_swaps", "成熟期日换电次数", "次/日",
              lambda c: sum(c.m("scale.mature_daily_swaps").values()),
              "Σ_池( 池内换电车辆 × 池内加权换电频次 )"),
            N("ops.onboard_avg", "全网加权装车电量（按换电次数）", "kWh",
              lambda c: (
                  c.m("swap_business.annual_energy_yi_kwh") * 1e8
                  / (sum(c.m("scale.mature_daily_swaps").values())
                     * c.p("swap_business.usable_energy_factor")
                     * c.p("swap_business.operating_days"))
              ),
              "由年电量反解的加权装车电量（各池装车电量按换电次数加权）"),
            P("ops.usable", "可用电量系数", "", "swap_business.usable_energy_factor"),
            P("ops.days", "年运营天数", "天", "swap_business.operating_days"),
        ],
    )

    rent_gwh = N(
        "ops.rent_gwh", "计费装机（可收租）", "GWh",
        lambda c: c.m("swap_business.rent_eligible_gwh"),
        "车端装机 ＋ 站内周转装机中归外部权益的部分",
        combine=lambda veh, ext: veh + ext,
        children=[
            N("ops.veh_gwh", "车端装机", "GWh",
              lambda c: c.m("swap_business.rent_vehicle_gwh"),
              "Σ_池( 池内换电车辆 × 池内加权装车电量 )"),
            N("ops.ext_gwh", "站内周转装机·外部权益部分", "GWh",
              lambda c: c.m("swap_business.rent_station_external_gwh"),
              "站内周转装机 × (1 − CATL建站持股比例)",
              combine=lambda st, own: st * (1.0 - own),
              children=[
                  N("ops.station_gwh", "站内周转装机", "GWh",
                    lambda c: sum(
                        c.m("capex.station_targets")[pk]
                        * c.p(f"stations.{pk}.inventory_blocks")
                        * c.p(f"stations.{pk}.block_kwh") / 1e6
                        for pk in POOLS),
                    "Σ_池( 站数 × 站内周转电池块数 × 单块电量 )"),
                  P("fin.ownership", "CATL建站持股比例", "", "finance.construction_ownership"),
              ]),
        ],
    )

    revenue = N(
        "ops.revenue", "换电业务年收入（成熟期）", "亿元",
        lambda c: c.m("swap_business.revenue_yi"),
        "服务费 ＋ 电池租金 ＋ 峰谷套利 ＋ 电网辅助服务",
        combine=lambda s, r, a, x: s + r + a + x,
        children=[
            N("ops.rev_service", "服务费收入", "亿元",
              lambda c: c.m("swap_business.service_revenue_yi"),
              "年换电交易电量 × 换电服务费单价",
              combine=lambda e, fee: e * fee,
              children=[energy,
                        P("ops.fee", "换电服务费单价", "元/kWh",
                          "swap_business.service_fee_rmb_kwh")]),
            N("ops.rev_rent", "电池租金收入", "亿元",
              lambda c: c.m("swap_business.battery_rent_yi"),
              "计费装机 × 单位装机年租金",
              combine=lambda g, rate: g * rate / 100.0,
              children=[rent_gwh,
                        P("ops.rent_rate", "单位装机年租金", "元/kWh·年",
                          "swap_business.battery_rent_rmb_kwh_year")]),
            N("ops.rev_arb", "峰谷套利收入", "亿元",
              lambda c: c.m("swap_business.arbitrage_yi"),
              "站内周转装机 × 运营天数 × 峰谷价差 × 站效率",
              combine=lambda g, days, spread, rte: g * days * spread * rte / 100.0,
              children=[
                  N("ops.station_gwh_ref", "站内周转装机", "GWh",
                    lambda c: sum(
                        c.m("capex.station_targets")[pk]
                        * c.p(f"stations.{pk}.inventory_blocks")
                        * c.p(f"stations.{pk}.block_kwh") / 1e6
                        for pk in POOLS),
                    "同上（引用）"),
                  P("ops.days_ref", "年运营天数", "天", "swap_business.operating_days"),
                  P("ops.spread", "峰谷价差", "元/kWh", "swap_business.grid_spread_rmb_kwh"),
                  P("ops.rte", "站内充放效率 RTE", "", "swap_business.rte"),
              ]),
            N("ops.rev_anc", "电网辅助服务收入", "亿元",
              lambda c: c.m("swap_business.ancillary_yi"),
              "容量补偿 ＋ 孰高( 需求响应, 调频 )　"
              "——容量补偿是站建成即可申报的「存在型」收益；后两项须实际提供服务才有收益，"
              "且争夺同一份可调容量，同一时段只能干一件事，故孰高取一、不可相加",
              combine=lambda a, b, d: a + max(b, d),
              children=[
                  N("ops.anc_cap", "容量补偿", "亿元",
                    lambda c: c.m("swap_business.capacity_compensation_yi"),
                    "Σ_池( 站数 × 单站申报容量 × 容量补偿价 )——站在即可申报，无条件计入"),
                  N("ops.anc_dr", "需求响应（候选·须实际提供服务）", "亿元",
                    lambda c: c.m("swap_business.demand_response_candidate_yi"),
                    "Σ_池( 站数 × 单站申报容量 × 年响应小时 × 响应价 )"),
                  N("ops.anc_fr", "调频（候选·须实际提供服务）", "亿元",
                    lambda c: c.m("swap_business.frequency_regulation_candidate_yi"),
                    "Σ_池( 站数 × 单站调频容量 × 年调频小时 × 调频价 )"),
              ]),
        ],
    )

    # ── 运营层：成本 ───────────────────────────
    opex = N(
        "ops.opex", "运营 OPEX", "亿元",
        lambda c: c.m("swap_business.opex_yi"),
        "电费 ＋ 场租 ＋ 人工 ＋ 软件调度 ＋ 保险 ＋ 池化维护 ＋ 仓储物流",
        combine=lambda e, r, l, s, i, p, w: e + r + l + s + i + p + w,
        children=[
            N("ops.cost_energy", "电费", "亿元",
              lambda c: c.m("swap_business.energy_cost_yi"),
              "( 充电电量 − 换电电量 ) × 谷电价",
              combine=lambda charged, out, price: (charged - out) * price,
              children=[
                  N("ops.charged_energy", "年充电电量", "亿kWh",
                    lambda c: (c.m("swap_business.annual_energy_yi_kwh")
                               / c.p("swap_business.rte")
                               * (1 + c.p("swap_business.auxiliary_power_rate"))),
                    "换电电量 ÷ RTE × ( 1 ＋ 站用电率 )",
                    combine=lambda e, rte, aux: e / rte * (1 + aux),
                    children=[
                        N("ops.energy_ref", "年换电交易电量", "亿kWh",
                          lambda c: c.m("swap_business.annual_energy_yi_kwh"), "同上（引用）"),
                        P("ops.rte_ref", "站内充放效率 RTE", "", "swap_business.rte"),
                        P("ops.aux", "站用电率", "", "swap_business.auxiliary_power_rate"),
                    ]),
                  N("ops.energy_ref2", "年换电交易电量", "亿kWh",
                    lambda c: c.m("swap_business.annual_energy_yi_kwh"), "同上（引用）"),
                  P("ops.valley_price", "谷段电价", "元/kWh",
                    "swap_business.valley_power_price_rmb_kwh"),
              ]),
            N("ops.cost_rent", "场地租金", "亿元",
              lambda c: c.m("swap_business.station_rent_yi"),
              "Σ_池( 站数 ) × 单站年场租",
              combine=lambda st, unit: st * unit / 1e4,
              children=[
                  N("st.total", "站数合计", "座",
                    lambda c: float(sum(c.m("capex.station_targets").values())),
                    "四个站型的终局站数之和"),
                  P("ops.site_rent", "单站年场租", "万元/年",
                    "swap_business.site_rent_wan_year"),
              ]),
            N("ops.cost_labor", "人工", "亿元",
              lambda c: c.m("swap_business.labor_yi"),
              "重卡站数 × 重卡单站年人工 ＋ 巧克力站数 × 巧克力单站年人工",
              combine=lambda hs, hw, cs, cw: (hs * hw + cs * cw) / 1e4,
              children=[
                  N("st.heavy", "重卡站数（终局）", "座",
                    lambda c: float(sum(n for k, n in c.m("capex.station_targets").items()
                                        if k.startswith("qiji75"))),
                    "骐骥短途站 ＋ 骐骥干线站"),
                  P("ops.labor_heavy", "重卡单站年人工", "万元/年",
                    "swap_business.heavy_station_labor_wan_year"),
                  N("st.choco", "巧克力站数（终局）", "座",
                    lambda c: float(sum(n for k, n in c.m("capex.station_targets").items()
                                        if k.startswith("choco"))),
                    "巧克力乘用站 ＋ 巧克力城配站"),
                  P("ops.labor_choco", "巧克力单站年人工", "万元/年",
                    "swap_business.passenger_station_labor_wan_year"),
              ]),
            N("ops.cost_software", "软件调度", "亿元",
              lambda c: c.m("swap_business.software_opex_yi"),
              "固定投入，按站数份额分池"),
            N("ops.cost_insurance", "保险", "亿元",
              lambda c: c.m("swap_business.insurance_yi"),
              "电池资产 × 电池保险费率 ＋ 站体资产 × 设备保险费率"),
            N("ops.cost_pooling", "池化维护", "亿元",
              lambda c: c.m("swap_business.pooling_maintenance_yi"),
              "电池资产 × 池化维护费率"),
            N("ops.cost_warehouse", "仓储物流", "亿元",
              lambda c: c.m("swap_business.warehouse_logistics_yi"),
              "电池资产 × 仓储物流费率"),
        ],
    )

    ebitda = N(
        "ops.ebitda", "EBITDA（成熟期）", "亿元",
        lambda c: c.m("swap_business.ebitda_yi"),
        "换电业务年收入 − 运营 OPEX",
        combine=lambda r, o: r - o,
        children=[revenue, opex],
    )

    # ── 资本层 ────────────────────────────────
    lifecycle = N(
        "cap.lifecycle", "全周期资本底座（现值）", "亿元",
        lambda c: c.m("capex.lifecycle_capital_base_yi"),
        "初装 CAPEX ＋ 全周期电池更新净额（按届时价格、扣回收残值，折现）",
        combine=lambda init, repl: init + repl,
        children=[
            N("cap.initial", "初装 CAPEX 合计", "亿元",
              lambda c: c.m("capex.total_initial_capex_yi"),
              "期初存量站投入 ＋ Σ_年( 车端电池 ＋ 站内电池 ＋ 站体 )",
              combine=lambda pre, veh, sb, body: pre + veh + sb + body,
              children=[
                  N("cap.preperiod", "期初存量站投入", "亿元",
                    lambda c: c.m("capex.preperiod_station_initial_capex_yi"),
                    "2026 年之前已建成站的站体与站内电池"),
                  N("cap.veh_battery", "车端电池初装", "亿元",
                    lambda c: sum(r["vehicle_battery_capex_yi"] for r in c.m("capex.annual")),
                    "Σ_年( 当年新增换电装车 GWh × 当年电池价格 )"),
                  N("cap.station_battery", "站内周转电池初装", "亿元",
                    lambda c: sum(r["station_battery_capex_yi"] for r in c.m("capex.annual")),
                    "Σ_年( 当年新增站数 × 站内周转块数 × 单块电量 × 当年电池价格 )"),
                  N("cap.station_body", "站体初装", "亿元",
                    lambda c: sum(r["station_body_capex_yi"] for r in c.m("capex.annual")),
                    "Σ_年Σ_池( 当年新增站数 × 该站型单站站体造价 )"),
              ]),
            N("cap.replacement", "全周期更新净额（现值）", "亿元",
              lambda c: (c.m("capex.lifecycle_capital_base_yi")
                         - c.m("capex.total_initial_capex_yi")),
              "全周期资本底座 − 初装 CAPEX（差额即更新净投入）"),
        ],
    )

    debt = N(
        "cap.debt", "项目债务", "亿元",
        lambda c: c.m("capex.project_debt_yi"),
        "全周期资本底座 × 债务比例",
        combine=lambda base, ratio: base * ratio,
        children=[lifecycle, P("fin.debt_ratio", "项目债务比例", "", "finance.debt_ratio")],
    )

    commitment = N(
        "cap.commitment", "CATL 全周期权益承诺", "亿元",
        lambda c: c.m("capex.catl_lifecycle_equity_commitment_yi"),
        "全周期资本底座 × (1 − 债务比例) × CATL建站持股比例",
        combine=lambda base, ratio, own: base * (1 - ratio) * own,
        children=[
            N("cap.lifecycle_ref", "全周期资本底座", "亿元",
              lambda c: c.m("capex.lifecycle_capital_base_yi"), "同上（引用）"),
            P("fin.debt_ratio_ref", "项目债务比例", "", "finance.debt_ratio"),
            P("fin.ownership_ref", "CATL建站持股比例", "", "finance.construction_ownership"),
        ],
    )

    capreq = N(
        "cap.annual_req", "年资本要求", "亿元",
        lambda c: c.m("capex.annual_capital_requirement_yi"),
        "全周期资本底座 × 资本回收系数 CRF",
        combine=lambda base, crf: base * crf,
        children=[
            N("cap.lifecycle_ref2", "全周期资本底座", "亿元",
              lambda c: c.m("capex.lifecycle_capital_base_yi"), "同上（引用）"),
            P("fin.crf", "资本回收系数 CRF", "", "finance.capital_recovery_factor"),
        ],
    )

    # ── 估值层：运营侧 ─────────────────────────
    ev = N(
        "val.ev", "运营侧企业价值 EV", "亿元",
        lambda c: c.m("swap_business.enterprise_value_yi"),
        "EBITDA × 运营 EV/EBITDA 倍数",
        combine=lambda e, mult: e * mult,
        children=[ebitda,
                  P("fin.ev_mult", "运营 EV/EBITDA 倍数", "×", "finance.swap_ev_ebitda")],
    )

    equity = N(
        "val.equity", "运营侧项目权益价值", "亿元",
        lambda c: c.m("swap_business.project_equity_value_yi"),
        "EV − 稳态debt（分池计算，单池权益不为负）",
        combine=lambda ev_, d: ev_ - d,
        children=[ev, N("cap.debt_ref", "稳态debt（2030在役资产历史成本×债务比）", "亿元",
                        lambda c: c.m("capex.steady_state_debt_yi"), "同上（引用）")],
    )

    op_value = N(
        "val.op_increment", "运营侧直接增量（CATL归属）", "亿元",
        lambda c: c.m("ledger.direct_swap_increment_value_yi"),
        "项目权益价值 × CATL建站持股比例",
        combine=lambda eq, own: eq * own,
        children=[equity,
                  P("fin.ownership_ref2", "CATL建站持股比例", "", "finance.construction_ownership")],
    )

    # ── 估值层：制造侧 ─────────────────────────
    mfg_np_gap = N(
        "val.mfg_np_gap", "制造净利差（有换电 − 纯制造）", "亿元",
        lambda c: c.m("ledger.full_manufacturing_scenario_gap_net_profit_yi"),
        "量差效应 ＋ 换电段利润率效应 ＋ 其他段利润率效应",
        combine=lambda vol, sm, om: vol + sm + om,
        children=[
            N("val.mfg_volume", "量差效应", "亿元",
              lambda c: c.m("ledger.manufacturing_volume_share_effect_net_profit_yi"),
              "锁定订单效应 ＋ 充电段份额效应（虹吸，通常为负）",
              combine=lambda lock, chg: lock + chg,
              children=[
                  N("val.mfg_lock", "锁定订单效应", "亿元",
                    lambda c: c.m("ledger.manufacturing_swap_locked_volume_effect_net_profit_yi"),
                    "换电锁定出货 × 单位售价 × 纯制造净利率"),
                  N("val.mfg_charge", "充电段份额效应（虹吸）", "亿元",
                    lambda c: c.m("ledger.manufacturing_charge_share_effect_net_profit_yi"),
                    "量差效应 − 锁定订单效应（换电对充电段的挤出）"),
              ]),
            N("val.mfg_margin_swap", "换电段利润率效应（护价）", "亿元",
              lambda c: c.m("ledger.manufacturing_swap_margin_effect_net_profit_yi"),
              "换电可覆盖收入 × 净利率差",
              combine=lambda rev, dm: rev * dm,
              children=[
                  N("val.mfg_swap_rev", "换电可覆盖收入", "亿元",
                    lambda c: c.m("ledger.with_swap_manufacturing.swap_addressable_revenue_yi")),
                  N("val.mfg_margin_delta", "净利率差（有换电 − 无换电）", "",
                    lambda c: (c.m("ledger.with_swap_manufacturing.net_margin")
                               - c.m("ledger.no_swap_manufacturing.net_margin")),
                    "有换电制造净利率 − 无换电制造净利率",
                    combine=lambda a, b: a - b,
                    children=[
                        N("val.margin_with", "有换电制造净利率", "",
                          lambda c: c.m("ledger.with_swap_manufacturing.net_margin")),
                        N("val.margin_no", "无换电制造净利率", "",
                          lambda c: c.m("ledger.no_swap_manufacturing.net_margin")),
                    ]),
              ]),
            N("val.mfg_margin_other", "其他段利润率效应", "亿元",
              lambda c: c.m("ledger.manufacturing_other_margin_effect_net_profit_yi"),
              "换电不可覆盖收入 × 净利率差"),
        ],
    )

    mfg_value = N(
        "val.mfg_increment", "制造侧锁量锁价增量", "亿元",
        lambda c: c.m("ledger.full_manufacturing_scenario_gap_value_yi"),
        "制造净利差 × 制造 PE",
        combine=lambda gap, pe: gap * pe,
        children=[mfg_np_gap, P("fin.mfg_pe", "制造 PE", "×", "finance.manufacturing_pe")],
    )

    increment = N(
        "val.increment", "可归因换电增量价值", "亿元",
        lambda c: c.m("ledger.total_swap_increment_value_yi"),
        "运营侧直接增量 ＋ 制造侧锁量锁价增量",
        combine=lambda op, mfg: op + mfg,
        children=[op_value, mfg_value],
    )

    # ── 体检层 ────────────────────────────────
    coverage = N(
        "chk.coverage", "EBITDA 覆盖倍数", "×",
        lambda c: c.m("swap_business.forward_to_required_ebitda"),
        "可交付 EBITDA ÷ 资本回报要求 EBITDA",
        combine=lambda e, req: e / req,
        children=[
            N("ops.ebitda_ref", "可交付 EBITDA", "亿元",
              lambda c: c.m("swap_business.ebitda_yi"), "同上（引用）"),
            N("chk.required_ebitda", "资本回报要求 EBITDA", "亿元",
              lambda c: c.m("swap_business.required_ebitda_yi"),
              "由年资本要求反推的门槛 EBITDA"),
        ],
    )

    ebit = N(
        "chk.ebit", "EBIT（账面）", "亿元",
        lambda c: c.m("swap_business.ebit_yi"),
        "EBITDA − 年折旧",
        combine=lambda e, d: e - d,
        children=[
            N("ops.ebitda_ref2", "EBITDA", "亿元",
              lambda c: c.m("swap_business.ebitda_yi"), "同上（引用）"),
            N("cap.depreciation", "年折旧（成熟期）", "亿元",
              lambda c: c.m("capex.mature_annual_depreciation_yi"),
              "电池折旧（按各池电池寿命） ＋ 站体折旧（站体总初装 ÷ 分析期）",
              combine=lambda batt, body, horizon: batt + body / horizon,
              children=[
                  N("cap.dep_battery", "电池折旧", "亿元",
                    lambda c: (c.m("capex.mature_annual_depreciation_yi")
                               - _body_total(c) / c.p("finance.model_horizon_years")),
                    "模型内按池 cohort 逐批计提：Σ_池( 池电池资本底座 ÷ 池电池寿命 )；"
                    "此处由成熟期折旧扣站体折旧反解"),
                  N("cap.body_total", "站体总初装（含期初存量站）", "亿元",
                    _body_total,
                    "Σ_年站体初装 ＋ 期初存量站数 × 该站型单站站体造价",
                    combine=lambda sched, pre: sched + pre,
                    children=[
                        N("cap.body_sched", "建设期站体初装", "亿元",
                          lambda c: sum(r["station_body_capex_yi"] for r in c.m("capex.annual")),
                          "Σ_年Σ_池( 当年新增站数 × 单站站体造价 )"),
                        N("cap.body_pre", "期初存量站站体", "亿元",
                          lambda c: sum(
                              c.m("capex.opening_station_stock_by_pool")[pk]
                              * c.p(f"stations.{pk}.station_body_capex_wan") / 1e4
                              for pk in c.m("capex.opening_station_stock_by_pool")),
                          "Σ_池( 期初存量站数 × 单站站体造价 )"),
                    ]),
                  P("fin.horizon", "模型分析期", "年", "finance.model_horizon_years"),
              ]),
        ],
    )

    dcf = N(
        "chk.dcf", "现金流支持的倍数（对照 拍的倍数）", "×",
        lambda c: c.m("swap_business.dcf_implied_multiple_at_crf"),
        "DCF企业价值（与门槛同源口径）÷ EBITDA　"
        "——回答「你用 EBITDA 的倍数估值，那现金流本身支持几倍」",
        combine=lambda ev_, e: ev_ / e,
        children=[
            N("chk.dcf_ev", "DCF企业价值·与门槛同源", "亿元",
              lambda c: c.m("swap_business.dcf_ev_at_crf_yi"),
              "成熟期 FCFF ÷ CRF（CRF 即设门槛 EBITDA 用的期望收益率年金因子，两者同源）",
              combine=lambda f, crf: f / crf,
              children=[
                  N("chk.fcff", "成熟期 FCFF", "亿元",
                    lambda c: c.m("swap_business.forward_fcff_yi"),
                    "EBITDA ×(1−税率) ＋ 折旧 × 税率（重置已由 CRF 年金化内含）"),
                  P("fin.crf_ref", "资本回收系数 CRF", "", "finance.capital_recovery_factor"),
              ]),
            N("ops.ebitda_ref3", "EBITDA", "亿元",
              lambda c: c.m("swap_business.ebitda_yi"), "同上（引用）"),
        ],
    )

    premium = N(
        "chk.premium", "战略溢价倍数（押注的大小）", "倍",
        lambda c: c.m("swap_business.dcf_multiple_premium"),
        "拍的 EV/EBITDA 倍数 ÷ 现金流支持的倍数　"
        "——这个数是判断不是计算，必须在报告正文里正面论证",
        combine=lambda mult, implied: mult / implied,
        children=[
            P("fin.ev_mult_ref", "运营 EV/EBITDA 倍数", "×", "finance.swap_ev_ebitda"),
            N("chk.dcf_ref", "现金流支持的倍数", "×",
              lambda c: c.m("swap_business.dcf_implied_multiple_at_crf"), "同上（引用）"),
        ],
    )

    ratio = N(
        "val.increment_pct", "增量价值 ÷ 集团市值", "",
        lambda c: c.m("ledger.attributable_swap_value_to_current_group_market_cap"),
        "可归因换电增量价值 ÷ 集团市值基准",
        combine=lambda inc, cap: inc / cap,
        children=[
            N("val.increment_ref", "可归因换电增量价值", "亿元",
              lambda c: c.m("ledger.total_swap_increment_value_yi"), "同上（引用）"),
            N("base.mktcap", "集团市值基准（A+H）", "亿元",
              lambda c: c.m("ledger.baseline.group_market_value_2026e_yi")),
        ],
    )

    # ── 估值层：NPV / 两条归属价值 / 押注 ──────────
    # 【新增 2026-09-03b】此前树只覆盖到「体检：过不过线」（覆盖倍数、隐含倍数），
    # 报告最终要用的三个数——NPV、CATL归属（DCF口径）、CATL归属（倍数法口径）——
    # 全都不在树里，这正是这次 debt 口径分裂能在两次 09-02 修复之后还漏到现在的原因。
    # 见 DECISIONS.md「2026-09-03b · 倍数法与DCF的debt口径分裂，暴露出一个更深的缺口」。
    valuation_capital_ref = N(
        "val.capital_pv", "估值资本PV（折到基年，DCF专用）", "亿元",
        lambda c: c.m("capex.valuation_capital_pv_yi"),
        "全部CAPEX统一折到base_year——与「做透的代价」分支的门槛口径 lifecycle_capital_base 不共用",
    )

    def _rebase_factor(c: "Ctx") -> float:
        return (1.0 + c.p("finance.wacc")) ** (
            c.p("meta.target_year") - c.p("construction.years")[0]
        )

    # 【新增 2026-09-04】把 base_year(2026) 的现值精确移到 target_year(2030)：
    # 对单一固定利率折现出的现值总额，乘 (1+wacc)^(target_year−base_year) 是精确
    # 操作、不是近似（推导见 capex_debt_估值公式链.md 第5.1-5.3节），树在这里
    # 自己重算一遍这个乘法，等于每次跑都在验证这条代数性质仍然成立。
    valuation_capital_target_ref = N(
        "val.capital_pv_target", "估值资本PV（移到target_year，DCF专用）", "亿元",
        lambda c: c.m("swap_business.dcf_valuation_capital_pv_at_target_yi"),
        "估值资本PV(base_year) × (1+WACC)^(target_year−base_year)——精确操作，非近似",
        combine=lambda base, factor: base * factor,
        children=[
            valuation_capital_ref,
            N("fin.rebase_factor", "移到target_year的复利因子", "", _rebase_factor,
              "(1+WACC)^(target_year−base_year)"),
        ],
    )

    terminal_residual_target_ref = N(
        "val.residual_pv_target", "期末残值现值（税后，移到target_year）", "亿元",
        lambda c: c.m("swap_business.dcf_terminal_residual_pv_at_target_yi"),
        "期末残值现值(base_year,税后) × (1+WACC)^(target_year−base_year)——精确操作，非近似",
        combine=lambda base, factor: base * factor,
        children=[
            N("val.residual_ref_base", "期末残值现值（税后，base_year）", "亿元",
              lambda c: c.m("swap_business.dcf_terminal_residual_pv_yi"),
              "各cohort期末残值按tax_rate计税、折回base_year后加总——"
              "见 DECISIONS「2026-09-03 · 末代残值不再降级储能」"),
            N("fin.rebase_factor2", "移到target_year的复利因子", "", _rebase_factor,
              "同上（引用）"),
        ],
    )

    # 【新增 2026-09-04】稳态debt本身按约58条cohort逐批计算（每条cohort取
    # target_year在役那一代的历史成本），不是两个既有节点的简单乘积，本节点
    # 不做递归自校验；一致性校验在 capex.py 构建时以「分池汇总=总量」的断言方式
    # 进行（steady_state_debt_base_by_pool 之和 == steady_state_debt_base）。
    dcf_debt = N(
        "val.dcf_debt", "稳态debt（2030在役资产历史成本×债务比，倍数法与DCF共用）", "亿元",
        lambda c: c.m("swap_business.dcf_valuation_debt_yi"),
        "Σ_cohort( target_year在役那一代的历史成本，含站体 ) × 债务比例——"
        "倍数法（val.catl_multiple）与DCF法在此共用同一个数，T3已解决",
    )

    npv = N(
        "val.npv", "项目NPV（DCF，CRF捷径EV，仅作对照）", "亿元",
        lambda c: c.m("swap_business.dcf_npv_at_crf_yi"),
        "DCF企业价值(CRF年金捷径) ＋ 期末残值现值(target_year) − 估值资本PV(target_year)",
        combine=lambda ev_, res, cap: ev_ + res - cap,
        children=[
            N("val.ev_ref", "DCF企业价值(CRF捷径)·与门槛同源", "亿元",
              lambda c: c.m("swap_business.dcf_ev_at_crf_yi"), "同上（引用）"),
            terminal_residual_target_ref,
            valuation_capital_target_ref,
        ],
    )

    catl_dcf = N(
        "val.catl_dcf", "CATL归属价值（DCF口径，CRF捷径EV，仅作对照）", "亿元",
        lambda c: c.m("swap_business.dcf_catl_value_at_crf_yi"),
        "max(0, DCF企业价值(CRF捷径) ＋ 期末残值现值(target_year) − 稳态debt) × 持股比例",
        combine=lambda ev_, res, d, own: max(0.0, ev_ + res - d) * own,
        children=[
            N("val.ev_ref2", "DCF企业价值(CRF捷径)", "亿元",
              lambda c: c.m("swap_business.dcf_ev_at_crf_yi"), "同上（引用）"),
            terminal_residual_target_ref,
            dcf_debt,
            P("fin.own_ref2", "CATL建站持股比例", "", "finance.construction_ownership"),
        ],
    )

    # 【2026-09-05，第四轮，见 capex_debt_估值公式链.md 第四轮 + DECISIONS
    # 「2026-09-04c」「2026-09-05」，本次去掉重复扣减】此前 ev_true 用
    # (forward_fcff − 折旧代理)×年金因子，而 npv_true/catl_dcf_true 又拿
    # valuation_capital_target（已经用真实排期精确扣过2031-2045更新支出）再减
    # 一次——同一笔钱被扣了两次。改回毛forward_fcff资本化，capex侧的真实排期
    # 已经在valuation_capital_target里扣过，不用再净一次。验证（wacc=7.5%）：
    # NPV从−1,371.9亿翻正到+1,650.9亿，隐含IRR从2.06%升到12.65%。
    ev_true = N(
        "val.ev_true", "DCF企业价值（有限期账/基础账，毛现金流资本化）", "亿元",
        lambda c: c.m("swap_business.dcf_ev_true_yi"),
        "成熟期FCFF（毛现金流，不扣任何资本性支出代理）× 15年WACC年金因子——"
        "资本性支出已经在 valuation_capital_target 里用真实排期精确扣过，这里"
        "不再重复扣减（此前用折旧代理再扣一次是bug，见 DECISIONS「2026-09-04c」）",
    )

    npv_true = N(
        "val.npv_true", "项目NPV（DCF，有限期账/基础账）", "亿元",
        lambda c: c.m("swap_business.dcf_npv_true_yi"),
        "DCF企业价值(毛现金流) ＋ 期末残值现值(电池,target_year) "
        "− 估值资本PV(target_year，含真实更新排期)。【2026-09-05d】站体设备不设"
        "期末残值——永续账已把它按'每15年到期全额换新、不扣回收'处理，隐含'到期"
        "视为耗尽'，有限期账不能再给同一台设备一份残值（两套相反假设），且没有"
        "设备二手市场数据支撑折价比例",
        combine=lambda ev_, res, cap: ev_ + res - cap,
        children=[ev_true, terminal_residual_target_ref, valuation_capital_target_ref],
    )

    catl_dcf_true = N(
        "val.catl_dcf_true", "CATL归属价值（DCF口径，有限期账/基础账，本轮新主口径）", "亿元",
        lambda c: c.m("swap_business.dcf_catl_value_true_yi"),
        "max(0, DCF企业价值(毛现金流) ＋ 期末残值现值(电池,target_year) "
        "− 稳态debt) × 持股比例（站体设备不设期末残值，理由同val.npv_true）",
        combine=lambda ev_, res, d, own: max(0.0, ev_ + res - d) * own,
        children=[
            ev_true, terminal_residual_target_ref, dcf_debt,
            P("fin.own_ref2b", "CATL建站持股比例", "", "finance.construction_ownership"),
        ],
    )

    # 【新增 2026-09-05，第四轮】永续账（开放上限）：不设15年截断、不含期末残值，
    # 可持续资本性支出改用 steady_state_net_replacement（更新理论第一性原理，非
    # 折旧代理）+ 递减永续年金（g=电池价格曲线长期降幅）+ 站体设备每horizon年一次
    # 性更新。见 capex_debt_估值公式链.md 第四轮§4.2-4.5。这三个节点不做递归自
    # 校验（分池renewal-theory anchor等中间量未逐一接入树，一致性校验已在
    # capex.py构建时以「分池汇总=总量」断言的方式进行）。
    ev_perpetual = N(
        "val.ev_perpetual", "DCF企业价值（永续账/开放上限）", "亿元",
        lambda c: c.m("swap_business.dcf_ev_perpetual_yi"),
        "成熟期FCFF÷WACC（永续，不设终点） − 稳态净更新支出÷(WACC+g)（递减永续，"
        "g=电池价格曲线长期降幅） − 站体设备永续更新PV（每horizon年一笔，标准"
        "递归年金公式 L÷[(1+WACC)^horizon−1]）——分池计算后加总，见第四轮§4.2-4.6",
    )

    npv_perpetual = N(
        "val.npv_perpetual", "项目NPV（DCF，永续账/开放上限）", "亿元",
        lambda c: c.m("swap_business.dcf_npv_perpetual_yi"),
        "DCF企业价值(永续) − 纯初装现值(target_year，剔除更换支出后的一次性投资，"
        "永续假设下不含期末残值——不会真的清算)",
    )

    catl_dcf_perpetual = N(
        "val.catl_dcf_perpetual", "CATL归属价值（DCF口径，永续账/开放上限）", "亿元",
        lambda c: c.m("swap_business.dcf_catl_value_perpetual_yi"),
        "max(0, DCF企业价值(永续) − 稳态debt) × 持股比例——与有限期账（基础账）"
        "共用同一个debt快照，差异只在EV怎么算（第一轮§1已论证debt该是资产负债表"
        "快照，跟期限假设无关）",
    )

    catl_multiple = N(
        "val.catl_multiple", "CATL归属价值（倍数法口径）", "亿元",
        lambda c: c.m("swap_business.catl_attributable_value_yi"),
        "max(0, EBITDA×拍定倍数 − 稳态debt) × 持股比例（分池计算，单池权益不为负；"
        "此处聚合层重算未做分池取0下限，与逐池加总可能有微小差异，属已知近似，"
        "见待办清单#12）",
        combine=lambda ebitda_, mult, d, own: max(0.0, ebitda_ * mult - d) * own,
        children=[
            N("ops.ebitda_ref4", "EBITDA", "亿元",
              lambda c: c.m("swap_business.ebitda_yi"), "同上（引用）"),
            P("fin.ev_mult_ref2", "运营EV/EBITDA倍数", "×", "finance.swap_ev_ebitda"),
            N("val.legacy_debt", "稳态debt（同DCF线，2026-09-04起统一）", "亿元",
              lambda c: c.m("capex.steady_state_debt_yi"),
              "2030年在役资产历史成本×债务比例——与DCF线共用同一个数（T3已解决）"),
            P("fin.own_ref3", "CATL建站持股比例", "", "finance.construction_ownership"),
        ],
    )

    debt_drift = N(
        "chk.debt_basis_drift", "稳态debt相对旧口径（门槛底座×债务比）改变了多少", "",
        lambda c: (
            (c.m("capex.steady_state_debt_yi") - c.m("swap_business.dcf_legacy_debt_yi"))
            / c.m("swap_business.dcf_legacy_debt_yi")
        ),
        "(新稳态debt − 旧门槛底座债务) ÷ 旧门槛底座债务——纯诊断：2026-09-04 起"
        "两条估值线已统一使用稳态debt，不再是「两条线口径分裂」（那个问题已解决），"
        "此节点只记录这次口径修正把debt改变了多少，见 DECISIONS「2026-09-03b」「2026-09-04」",
    )

    framework_ul_gap = N(
        "chk.framework_ul_gap", "框架L(折现FCFE)相对框架U(EV−debt)的偏离", "%",
        lambda c: c.m("swap_business.dcf_framework_ul_gap_pct"),
        "【T5，2026-09-04，纯诊断】(框架L股权价值 − 框架U股权价值) ÷ 框架U股权价值。"
        "框架L：Ke=[WACC−债务比×债务利率×(1−税率)]÷(1−债务比) 折现FCFE"
        "（FCFE=真实FCFF−税后利息，debt按steady_state_debt_yi固定不摊销、"
        "净新增借款=0，与forward_fcff/sustaining_capex同样的'稳态不变'假设）。"
        "两者不精确相等是预期内的：附录'情形C'证明U=L恒等式要求debt按"
        "'债务比×剩余项目价值'逐期摊销到0（value降、debt跟着降）；本模型debt是"
        "'target_year在役资产历史成本'快照，15年窗口内固定不变——这是两种不同的"
        "'稳定杠杆'定义，有限期+终值截断场景下不再代数等价。此gap量化这一口径差异"
        "的幅度，不参与上层校验、不代表程序错误",
    )

    bet = N(
        "val.bet", "押注部分（结构性差值）", "亿元",
        lambda c: c.m("swap_business.dcf_catl_value_gap_yi"),
        "倍数法归属 − DCF归属(真实口径)　"
        "【2026-09-04：debt统一后，两条线现在真的在问同一个问题、用同一个debt，"
        "此前藏着的第三份「倍数法归属」（_dcf_cross_check内部现算、从未暴露的版本，"
        "见 DECISIONS「2026-09-03b」）已删除，三份合并为一份（T3已解决），此处直接用"
        "真正的倍数法headline数字相减】",
        combine=lambda m, d: m - d,
        children=[catl_multiple, catl_dcf_true],
    )

    valuation = N(
        "branch.valuation", "估值：这门生意最终值多少钱", "亿元",
        lambda c: c.m("swap_business.catl_attributable_value_yi"),
        "见下——倍数法与DCF两套独立口径（DCF区分CRF捷径对照版 npv/catl_dcf、"
        "有限期账/基础账 npv_true/catl_dcf_true、永续账/开放上限"
        "npv_perpetual/catl_dcf_perpetual——两本账都不建模2030后规模增长"
        "（capex诊断.md§8.4已论证的既定假设），差异只在'停在2030规模不动'延续"
        "多久：15年还是永续，见第四轮§4.7），「押注」是倍数法与DCF基础账的结构性"
        "差值；debt_basis_drift、framework_ul_gap 只诊断、不参与本节点校验",
        combine=(
            lambda npv_, catl_dcf_, npv_t_, catl_dcf_t_, npv_p_, catl_dcf_p_,
            catl_mult_, bet_, drift_, ul_gap_: catl_mult_
        ),
        children=[
            npv, catl_dcf, npv_true, catl_dcf_true, npv_perpetual, catl_dcf_perpetual,
            catl_multiple, bet, debt_drift, framework_ul_gap,
        ],
    )

    # ── 根 ────────────────────────────────────
    return N(
        "root", "换电这笔投资值不值得做", "",
        lambda c: 0.0,
        "由「换回的价值 / 做透的代价 / 体检过不过线 / 估值最终值多少」四支共同回答，不是一个数",
        children=[
            N("branch.value", "换回的价值", "亿元",
              lambda c: c.m("ledger.total_swap_increment_value_yi"),
              "见下",
              combine=lambda inc, pct: inc,
              children=[increment, ratio]),
            N("branch.cost", "做透的代价", "亿元",
              lambda c: c.m("capex.catl_lifecycle_equity_commitment_yi"),
              "见下",
              combine=lambda commit, base, req, d: commit,
              children=[commitment, lifecycle, capreq, debt]),
            N("branch.check", "体检：过不过线", "",
              lambda c: c.m("swap_business.forward_to_required_ebitda"),
              "见下",
              combine=lambda cov, ebit_, dcf_, prem: cov,
              children=[coverage, ebit, dcf, premium]),
            valuation,
        ],
    )


# ─────────────────────────────────────────── 校验
@dataclass
class Finding:
    kind: str        # mismatch | no_source | no_formula
    key: str
    label: str
    detail: str


def walk(node: Node):
    yield node
    for child in node.children:
        yield from walk(child)


def verify(root: Node, c: Ctx) -> list[Finding]:
    out: list[Finding] = []
    for node in walk(root):
        # ① 父 = f(子)
        if node.combine is not None and node.children:
            try:
                got = float(node.value(c))
                want = float(node.combine(*[float(ch.value(c)) for ch in node.children]))
            except Exception as exc:  # noqa: BLE001
                out.append(Finding("mismatch", node.key, node.label, f"取值/重算抛错：{exc!r}"))
                continue
            scale = max(abs(got), abs(want), 1e-12)
            rel = abs(got - want) / scale
            if rel > node.tol:
                out.append(Finding(
                    "mismatch", node.key, node.label,
                    f"模型值 {got:,.6g} vs 按公式重算 {want:,.6g}"
                    f"（相对差 {rel:.3%}，容差 {node.tol:.1e}）"
                    + (f"｜放宽理由：{node.tol_note}" if node.tol_note else "")
                ))
        # ② 叶节点必须有信源
        if node.is_leaf and node.param:
            if not node.source and not c.src_of(node.param):
                out.append(Finding("no_source", node.key, node.label,
                                   f"base.toml 里 `{node.param}` 上方没有说明取值依据的注释"))
        # ③ 非叶节点必须有公式
        if not node.is_leaf and not node.formula:
            out.append(Finding("no_formula", node.key, node.label, "缺公式说明"))
    return out


# ─────────────────────────────────────────── 导出
def to_dict(node: Node, c: Ctx, scen_ctx: dict[str, Ctx] | None = None) -> dict:
    doc = c.docs.get(node.param or "", {})
    src = c.src_of(node.param) if node.param else None

    def _v(ctx: Ctx):
        try:
            return float(node.value(ctx))
        except Exception:  # noqa: BLE001
            return None

    out = {
        "key": node.key,
        "label": node.label,
        "unit": node.unit,
        "value": _v(c),
        "formula": node.formula,
        "param": node.param,
        "source": node.source or src or None,
        "section": doc.get("section") or None,
    }
    if scen_ctx:
        out["scenarios"] = {name: _v(ctx) for name, ctx in scen_ctx.items()}
    out["children"] = [to_dict(ch, c, scen_ctx) for ch in node.children]
    return out


def export_xlsx(root: Node, c: Ctx, scen_ctx: dict[str, Ctx], path: Path) -> None:
    """写出可折叠的层级表：Excel 分级显示（左侧 +/− 展开），列＝三情景。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "决策树"
    headers = ["层", "节点", *[n for n in scen_ctx], "单位", "计算链（公式）",
               "参数路径（叶节点）", "取值依据 / 信源"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor="EFF0F3")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    def emit(node: Node, depth: int) -> None:
        doc = c.docs.get(node.param or "", {})
        src = (node.source or (c.src_of(node.param) if node.param else "")
               or doc.get("section", "") or "")
        vals = []
        for ctx in scen_ctx.values():
            try:
                vals.append(round(float(node.value(ctx)), 4))
            except Exception:  # noqa: BLE001
                vals.append(None)
        ws.append([depth, "　" * depth + node.label, *vals, node.unit,
                   node.formula, node.param or "", src])
        row = ws.max_row
        if depth:
            ws.row_dimensions[row].outlineLevel = min(depth, 7)
            ws.row_dimensions[row].hidden = depth > 2      # 默认只展开两层
        if node.is_leaf:
            for col in (2, len(headers)):
                ws.cell(row=row, column=col).font = Font(italic=True, size=9)
        for ch in node.children:
            emit(ch, depth + 1)

    emit(root, 0)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = "C2"
    widths = [5, 42, *[13] * len(scen_ctx), 9, 58, 40, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


# ─────────────────────────────────────────── 主流程
def main() -> None:
    parser = argparse.ArgumentParser(description="以终为始决策树 · 链路审计")
    parser.add_argument("--json", action="store_true", help="额外写出 outputs/tree.json")
    parser.add_argument("--xlsx", action="store_true", help="额外写出可折叠的 Excel 层级表")
    args = parser.parse_args()

    scen_ctx = {name: build_ctx(scen) for name, scen in SCENARIOS}
    c = scen_ctx["中性"]          # 中性档严格等于基线，审计以它为准

    root = build_tree(c)
    nodes = list(walk(root))
    leaves = [n for n in nodes if n.is_leaf]
    checked = [n for n in nodes if n.combine is not None and n.children]

    print("═" * 64)
    print("以终为始决策树 · 链路审计")
    print("═" * 64)
    print(f"节点 {len(nodes)}　叶节点 {len(leaves)}"
          f"（其中直达 base.toml 参数 {sum(1 for n in leaves if n.param)}）"
          f"　带公式可校验的内部节点 {len(checked)}")
    print(f"情景：{'／'.join(scen_ctx)}（中性＝模型基线）")
    print()

    titles = {
        "mismatch": "① 公式与模型对不上（父 ≠ f(子)）",
        "no_source": "② 叶参数缺信源（base.toml 里没有取值依据注释）",
        "no_formula": "③ 内部节点缺公式说明",
    }
    total = 0
    for scen_name, ctx in scen_ctx.items():
        findings = verify(build_tree(ctx), ctx)
        total += len(findings)
        if not findings:
            print(f"【{scen_name}】全部通过")
            continue
        print(f"【{scen_name}】{len(findings)} 处")
        by_kind: dict[str, list[Finding]] = {}
        for f in findings:
            by_kind.setdefault(f.kind, []).append(f)
        for kind in ("mismatch", "no_source", "no_formula"):
            for f in by_kind.get(kind, []):
                print(f"   {titles[kind][:2]} [{f.key}] {f.label}")
                print(f"       {f.detail}")
    print()
    if not total:
        print("三个情景下，树的每个内部节点都自洽，每个叶参数都有信源。")

    if args.json:
        TREE_PATH.write_text(
            json.dumps(to_dict(root, c, scen_ctx), ensure_ascii=False, indent=1) + "\n",
            encoding="utf-8")
        print(f"已写出 {TREE_PATH.name}")
    if args.xlsx:
        export_xlsx(root, c, scen_ctx, XLSX_PATH)
        print(f"已写出 {XLSX_PATH.name}")


if __name__ == "__main__":
    main()
